"""Build slide decks the way people actually export them, for the corpus.

  slides-libreoffice.pdf  a PowerPoint deck (python-pptx) exported by LibreOffice Impress
  slides-chromium.pdf     a web slide deck printed by Chromium -- the Skia engine Google Slides uses

Both carry what slide exports are known to break PDF/A pipelines with: 16:9 pages, gradients,
transparency and soft masks, photos with alpha, rotated text, charts and tables as vectors, blend
modes, shadows, symbols and colour emoji. Runs in the image's `slides` build stage; on a Mac it uses
LibreOffice and Google Chrome if installed (SOFFICE / CHROME override the paths).

    python make_slides.py OUT_DIR
"""

from __future__ import annotations

import base64
import io
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

TEXT = [
    "Umsatz Q3 2026: 1.284.500 EUR (+12,4 %)",
    "Neue Kunden in München, Köln und İstanbul",
    "Rückstellungen für Gewährleistung: 35.000 EUR",
    "Nächste Schritte: Prüfung durch Frau Öztürk",
]


def photo(alpha: bool = True) -> bytes:
    """A synthetic 'photo': gradient sky, sun disc, translucent where alpha is on."""
    image = Image.new("RGBA", (800, 500))
    draw = ImageDraw.Draw(image)
    for y in range(500):
        draw.line([(0, y), (800, y)], fill=(40 + y // 5, 110 + y // 8, 210 - y // 4, 255 if not alpha else 230))
    draw.ellipse((260, 90, 540, 370), fill=(255, 205, 40, 190 if alpha else 255))
    draw.rectangle((0, 400, 800, 500), fill=(40, 120, 60, 255))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def set_alpha(shape, percent: int) -> None:
    """python-pptx has no API for fill transparency; add <a:alpha> to the solid fill colour."""
    fill = shape.fill._xPr.find(qn("a:solidFill"))
    colour = fill[0]
    alpha = colour.makeelement(qn("a:alpha"), {"val": str(percent * 1000)})
    colour.append(alpha)


def build_pptx(path: Path, image: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.gradient()
    background.fill.gradient_angle = 35
    background.fill.gradient_stops[0].color.rgb = RGBColor(0x0B, 0x3D, 0x91)
    background.fill.gradient_stops[1].color.rgb = RGBColor(0x55, 0xC1, 0xE8)
    background.line.fill.background()
    title = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.6)).text_frame
    title.text = "Quartalsbericht Müller & Şahin GmbH"
    title.paragraphs[0].runs[0].font.size = Pt(48)
    title.paragraphs[0].runs[0].font.bold = True
    title.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    subtitle = slide.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(11.5), Inches(0.8)).text_frame
    subtitle.text = "Stuttgart · 12.03.2026 · vertraulich"
    subtitle.paragraphs[0].runs[0].font.size = Pt(24)
    subtitle.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xE8, 0xF4, 0xFF)

    slide = prs.slides.add_slide(blank)
    heading = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1)).text_frame
    heading.text = "Ergebnisse im Überblick"
    heading.paragraphs[0].runs[0].font.size = Pt(36)
    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(6.8), Inches(4.5)).text_frame
    body.word_wrap = True
    for index, line in enumerate(TEXT):
        paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
        paragraph.text = "• " + line
        paragraph.runs[0].font.size = Pt(22)
    slide.shapes.add_picture(str(image), Inches(7.8), Inches(1.6), Inches(5), Inches(3.1))
    veil = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(4.3), Inches(5.6), Inches(1.6))
    veil.fill.solid()
    veil.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    set_alpha(veil, 55)
    veil.text_frame.text = "Bildquelle: eigene Darstellung"
    stamp = slide.shapes.add_textbox(Inches(8.5), Inches(5.9), Inches(4), Inches(1)).text_frame
    stamp.text = "ENTWURF"
    stamp.paragraphs[0].runs[0].font.size = Pt(40)
    stamp.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xD0, 0x20, 0x20)
    slide.shapes[-1].rotation = -18

    slide = prs.slides.add_slide(blank)
    heading = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1)).text_frame
    heading.text = "Umsatz nach Region (Tsd. EUR)"
    heading.paragraphs[0].runs[0].font.size = Pt(32)
    data = CategoryChartData()
    data.categories = ["Nord", "Süd", "Ost", "West"]
    data.add_series("2025", (310, 420, 180, 250))
    data.add_series("2026", (345, 468, 202, 269))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.5), Inches(6.5), Inches(5), data)
    table = slide.shapes.add_table(5, 3, Inches(7.6), Inches(1.7), Inches(5.2), Inches(3)).table
    for row, values in enumerate(
        [
            ("Region", "2025", "2026"),
            ("Nord", "310", "345"),
            ("Süd", "420", "468"),
            ("Ost", "180", "202"),
            ("West", "250", "269"),
        ]
    ):
        for column, value in enumerate(values):
            table.cell(row, column).text = value

    slide = prs.slides.add_slide(blank)
    heading = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1)).text_frame
    heading.text = "Status der Maßnahmen ✓ ✗ → 📈"
    heading.paragraphs[0].runs[0].font.size = Pt(32)
    for index, (label, colour) in enumerate([("Erledigt ✓", "2E7D32"), ("Offen ✗", "C62828"), ("Läuft →", "F9A825")]):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + index * 4.1), Inches(2.2), Inches(3.6), Inches(2.2)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(colour)
        set_alpha(box, 80)
        box.text_frame.text = label
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(0.8), Inches(5.2), Inches(11.8), Inches(0.9))
    arrow.fill.gradient()
    arrow.fill.gradient_stops[0].color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    arrow.fill.gradient_stops[1].color.rgb = RGBColor(0x0B, 0x3D, 0x91)
    prs.save(path)


HTML = """<!doctype html><html lang="de"><head><meta charset="utf-8"><style>
@page {{ size: 1280px 720px; margin: 0 }}
body {{ margin: 0; font-family: 'DejaVu Sans', 'Liberation Sans', Arial, sans-serif }}
.slide {{ width: 1280px; height: 720px; position: relative; overflow: hidden; page-break-after: always; box-sizing: border-box; padding: 60px 80px }}
.title {{ background: linear-gradient(135deg, #0b3d91 0%, #55c1e8 60%, #ffd35c 100%); color: white }}
h1 {{ font-size: 56px; margin: 180px 0 10px }} h2 {{ font-size: 40px; margin: 0 0 30px }}
.card {{ position: absolute; border-radius: 24px; box-shadow: 0 18px 40px rgba(0,0,0,.35); padding: 24px 32px; font-size: 26px }}
.glass {{ background: rgba(255,255,255,.55); backdrop-filter: blur(6px) }}
.photo {{ position: absolute; right: 80px; top: 150px; width: 520px; height: 330px; border-radius: 40px; object-fit: cover; opacity: .9 }}
.blend {{ mix-blend-mode: multiply; background: radial-gradient(circle, #ff6b6b 0%, rgba(255,107,107,0) 70%);
          position: absolute; width: 500px; height: 500px; right: -120px; bottom: -160px }}
.rotated {{ position: absolute; right: 120px; bottom: 60px; font-size: 44px; color: #d02020; transform: rotate(-18deg); opacity: .7 }}
li {{ font-size: 28px; margin: 12px 0 }}
</style></head><body>
<section class="slide title"><h1>Quartalsbericht Müller &amp; Şahin GmbH</h1><p style="font-size:28px">Stuttgart · 12.03.2026 · vertraulich</p>
<div class="blend"></div></section>
<section class="slide"><h2>Ergebnisse im Überblick 📈</h2><ul>{items}</ul>
<img class="photo" src="data:image/png;base64,{photo}"><div class="card glass" style="right:110px; top:420px">Bildquelle: eigene Darstellung</div>
<div class="rotated">ENTWURF</div></section>
<section class="slide"><h2>Umsatz nach Region (Tsd. EUR)</h2>
<svg width="620" height="420" viewBox="0 0 620 420">
  <defs><linearGradient id="g" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#55c1e8"/><stop offset="1" stop-color="#0b3d91"/></linearGradient></defs>
  {bars}
  <line x1="40" y1="380" x2="600" y2="380" stroke="#333" stroke-width="2"/>
</svg>
<div class="card" style="left:760px; top:170px; background:#f4f7fb">Nord 345 · Süd 468<br>Ost 202 · West 269<br><b>Summe 1.284</b> ✓</div>
</section></body></html>"""


def build_html(path: Path, image: bytes) -> None:
    values = [(310, 345), (420, 468), (180, 202), (250, 269)]
    bars = []
    for index, (before, after) in enumerate(values):
        x = 60 + index * 140
        bars.append(
            f'<rect x="{x}" y="{380 - before * 0.7:.0f}" width="50" height="{before * 0.7:.0f}" fill="#b0bec5" opacity=".8"/>'
        )
        bars.append(
            f'<rect x="{x + 56}" y="{380 - after * 0.7:.0f}" width="50" height="{after * 0.7:.0f}" fill="url(#g)"/>'
        )
        bars.append(f'<text x="{x + 20}" y="405" font-size="18">{["Nord", "Süd", "Ost", "West"][index]}</text>')
    path.write_text(
        HTML.format(
            items="".join(f"<li>{line}</li>" for line in TEXT),
            photo=base64.b64encode(image).decode(),
            bars="".join(bars),
        ),
        encoding="utf-8",
    )


def find(env: str, *candidates: str) -> str:
    for candidate in (os.environ.get(env), *candidates):
        if candidate and (shutil.which(candidate) or Path(candidate).exists()):
            return shutil.which(candidate) or candidate
    sys.exit(f"missing {env}: none of {candidates} found")


def print_pdf(chrome: str, work: Path, url: str, target: Path, timeout: float = 180) -> None:
    """Chrome writes the PDF and, on macOS, then keeps running; wait for a settled file, then end it."""
    target.unlink(missing_ok=True)
    process = subprocess.Popen(
        [
            chrome,
            "--headless",
            "--no-sandbox",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--user-data-dir={work}/chrome-profile",
            f"--print-to-pdf={target}",
            url,
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    deadline, last = time.monotonic() + timeout, -1
    try:
        while time.monotonic() < deadline:
            if process.poll() is not None and target.exists():
                return
            size = target.stat().st_size if target.exists() else -1
            if size > 0 and size == last:
                return
            last = size
            time.sleep(1)
        sys.exit("Chrome did not print the deck in time")
    finally:
        if process.poll() is None:
            process.terminate()
            try:
                process.wait(10)
            except subprocess.TimeoutExpired:
                process.kill()


def main() -> None:
    out = Path(sys.argv[1]).resolve()
    out.mkdir(parents=True, exist_ok=True)
    soffice = find("SOFFICE", "soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice")
    chrome = find(
        "CHROME", "chromium", "chromium-browser", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
    )
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        picture = work / "photo.png"
        picture.write_bytes(photo())
        build_pptx(work / "slides-libreoffice.pptx", picture)
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation=file://{work}/lo-profile",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(out),
                str(work / "slides-libreoffice.pptx"),
            ],
            check=True,
            capture_output=True,
        )
        build_html(work / "deck.html", photo())
        print_pdf(chrome, work, (work / "deck.html").as_uri(), out / "slides-chromium.pdf")
    for name in ("slides-libreoffice.pdf", "slides-chromium.pdf"):
        if not (out / name).exists():
            sys.exit(f"{name} was not produced")
        print(f"{name}: {(out / name).stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
